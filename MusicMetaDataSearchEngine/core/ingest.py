"""Metadaten-Eingang: Library-Daten (CSV/XLSX/XLS-Dumps) -> songs.db einspeisen.

Manche Labels/Libraries schicken regelmäßig ihre Datenbank als Datei (keine API).
Dieser Importer erkennt die Spalten automatisch (DE+EN-Aliasse) und merged per
upsert in die wachsende Song-DB, ohne schon verifizierte Werte zu überschreiben.

Formate: .csv (Trenner wird geschnüffelt: , ; oder Tab), .xlsx (openpyxl),
.xls (xlrd, falls installiert). Neben den Kernfeldern werden auch die
GEMA-schlanken Zusatzfelder (album, track_number, isrc, iswc, gema_werknummer)
erkannt und gemerged.
"""
from __future__ import annotations
import csv
import io
import re
from pathlib import Path

from . import songsdb

# unser Feld -> mögliche Spaltennamen in fremden Dumps. Verglichen wird normalisiert (klein, ohne alles
# außer Buchstaben und Ziffern), deshalb deckt ein Eintrag alle Schreibweisen ab: „LC code", „LC-Code",
# „lc_code" sind dasselbe. Innerhalb einer Liste gewinnt der erste Treffer — Spezifisches nach vorn.
_ALIASES = {
    "interpret": ["interpret", "artist", "artists", "kuenstler", "künstler", "performer", "performed by",
                  "interpreten", "ausführender künstler"],
    "titel": ["titel", "title", "track title", "work title", "song title", "cue title", "musiktitel",
              "trackname", "songtitle", "track", "song", "werktitel"],
    "komponist": ["komponist", "komponisten", "composer", "composers", "composer/author", "urheber",
                  "music", "musik"],
    "texter": ["texter", "textdichter", "lyricist", "lyricists", "lyrics by", "author", "autor"],
    "verlag": ["verlag", "musikverlag", "original publisher", "originalverlag", "publisher", "publishers",
               "publisher name", "music publisher", "verleger"],
    "labelname": ["labelname", "label name", "label", "plattenlabel", "plattenfirma", "labelname_gvl"],
    "labelcode": ["labelcode", "label code", "lc code", "german label code", "gvl code", "gvl lc", "lc", "gvl"],
    "album": ["album", "album title", "albumtitel", "release title", "release", "albumname"],
    "track_number": ["track number", "track no", "tracknr", "track nr", "titelnummer", "cd track"],
    "isrc": ["isrc", "isrc code"],
    "iswc": ["iswc", "iswc code", "werk iswc"],
    # „GEMA work no.“: so heißt das Feld in der englischen Oberfläche — ohne eigenen Alias griff es nur über den
    # Teilstring „gema“, und eine andere GEMA-Spalte weiter vorn im Kopf hätte es verdrängt
    "gema_werknummer": ["gema werknummer", "gema werknr", "gema work code", "gema work number", "gema work no",
                        "code gema work number", "werknummer", "gema werk", "gema"],
}

# Nur diese Felder dürfen im zweiten Durchgang als TEIL einer zusammengesetzten Kopfzeile greifen
# („CODE: GEMA WORK NUMBER" -> gema_werknummer). Für die übrigen wäre das zu grob: „Sub Label" ist
# kein Labelname, „Track ID" keine Track-Nummer.
_SUBSTRING_OK = ("gema_werknummer", "labelcode", "isrc", "iswc", "verlag", "komponist", "titel")
# Als Teilstring zählen Aliasse ab vier Zeichen („lc", „gvl" träfen zu viel: „Calculated"). Ausgenommen sind
# Allerweltswörter, die in fremden Köpfen stecken: „track" in „Soundtrack Album", „song" in „Songwriter",
# „music" in „Music Library". Eine pauschale Grenze von sechs Zeichen warf auch „ISRC-Nummer", „Work ISWC",
# „Titel des Werkes" und „GEMA-Nr." raus.
_SUBSTRING_TOO_GENERIC = frozenset({"track", "song", "music", "musik"})


def _norm(name: str) -> str:
    """Spaltenname -> Vergleichsform: klein, ohne Leerzeichen, Punkte, Doppelpunkte, Bindestriche."""
    return re.sub(r"[^a-z0-9]+", "", (name or "").strip().lower())


SUPPORTED_EXTENSIONS = (".csv", ".xlsx", ".xls")
FIELDS = tuple(_ALIASES)          # unsere Zielfelder in Anzeige-Reihenfolge


def _cell(v) -> str:
    """Zellwert -> getrimmter String. Ganzzahlige Floats (openpyxl/xlrd liefern
    Zahlen als float) werden ohne „.0" ausgegeben, damit z.B. track_number 2.0 -> „2"."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _decode(path: str | Path) -> str:
    """Dateiinhalt als Text. UTF-16 (BOM) kommt von Windows-Exporten — Cue-Sheets mancher Libraries
    sind so kodiert; ohne Erkennung liest der Parser eine einzige Spalte voll Nullbytes."""
    raw = Path(path).read_bytes()
    for bom, enc in ((b"\xff\xfe\x00\x00", "utf-32"), (b"\x00\x00\xfe\xff", "utf-32"),
                     (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16"), (b"\xef\xbb\xbf", "utf-8-sig")):
        if raw.startswith(bom):
            return raw.decode(enc)
    for enc in ("utf-8", "cp1252", "latin-1"):      # latin-1 kann alles, ist also immer die letzte Rettung
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _header_index(grid: list[list[str]]) -> int:
    """Zeilenindex der Kopfzeile. GEMA-Cue-Sheets und manche Exporte haben Titelblock und Leerzeilen
    darüber (';;;Music Cue Sheet;;;;' — die Kopfzeile steht dort erst in Zeile 14). Gewählt wird unter
    den ersten 20 Zeilen die mit den meisten gefüllten Zellen; bei Gleichstand die oberste.

    Mit Toleranz: eine Datenzeile, die zufällig eine Zelle mehr gefüllt hat als die Kopfzeile (etwa
    weil deren letzte Spalte unbenannt ist), darf die Kopfzeile nicht verdrängen. Deshalb gewinnt die
    OBERSTE Zeile, die bis auf eine Zelle an das Maximum herankommt.

    Die Toleranz allein reicht bei schmalen Listen nicht: bei drei Spalten ist ein Vorspann wie
    'Projekt;ZZQ Testshow 1234' genau Maximum − 1 breit und gewann. Unter den Kandidaten
    gewinnt deshalb die Zeile mit den meisten erkannten Spaltennamen; erst bei Gleichstand die oberste."""
    filled = [sum(1 for c in row if (c or "").strip()) for row in grid[:20]]
    if not filled:
        return 0
    best_n = max(filled)
    if best_n < 3:
        return 0
    cands = [i for i, n in enumerate(filled) if n >= best_n - 1]
    return max(cands, key=lambda i: (len(guess_mapping(grid[i])), -i))


def _dialect(text: str, forced: str = ""):
    """Trenner bestimmen. csv.Sniffer riet bei TRACKLISTE.csv auf Komma, weil die Komponisten-Spalte
    „Vorname Nachname, Vorname Nachname" enthält — die ganze Zeile landete in einer Spalte.
    Deshalb gezählt statt geraten: welcher Kandidat zerlegt die ersten Zeilen am gleichmäßigsten?"""
    if forced:
        return type("_d", (csv.excel,), {"delimiter": forced})
    probe = text[:16384]
    best, best_score = ",", (0.0, 0)
    for d in (",", ";", "\t", "|"):
        rows = [r for r in list(csv.reader(io.StringIO(probe), delimiter=d))[:20] if r]
        if len(rows) < 2:
            continue
        counts = [len(r) for r in rows]
        common = max(set(counts), key=counts.count)      # häufigste Spaltenzahl = die echte Breite
        if common < 2:
            continue
        # Erst Gleichmäßigkeit, dann Breite: das Komma zerlegt TRACKLISTE.csv zwar ebenso gleichmäßig
        # (die Komponisten-Spalte hat genau ein Komma), aber nur in zwei Spalten statt in zehn.
        score = (round(counts.count(common) / len(counts), 2), common)
        if score > best_score:
            best, best_score = d, score
    return type("_d", (csv.excel,), {"delimiter": best})


def _read_csv(path: str | Path) -> tuple[list[str], list[dict]]:
    text = _decode(path)
    # Excel schreibt als erste Zeile gern 'sep=,': die ist keine
    # Kopfzeile, aber sie nennt den Trenner. Ohne diesen Sprung liest der Parser 'sep=' als einzige
    # Spalte und meldet, es fehlten alle erwarteten Spalten.
    forced = ""
    head, nl, rest = text.partition("\n")
    if head.strip().lower().startswith("sep=") and len(head.strip()) >= 5:
        forced, text = head.strip()[4:5], rest
    grid = list(csv.reader(io.StringIO(text), _dialect(text, forced)))
    if not grid:
        return [], []
    h = _header_index(grid)
    header, seen = [], {}
    for i, name in enumerate(grid[h]):
        name = (name or "").strip() or f"Spalte {i + 1}"
        if name in seen:                      # doppelte Köpfe (GEMA-Bogen hat 'Share %' zweimal)
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        header.append(name)
    rows = []
    for row in grid[h + 1:]:
        if not row:
            continue      # echte Leerzeile — eine Zeile aus lauter leeren Zellen (';;;') bleibt drin,
                          # sonst verschieben sich die Zeilennummern in den Warnungen (librarylists)
        rows.append({header[i]: (c or "").strip() for i, c in enumerate(row) if i < len(header)})
    return header, rows


def _read_xlsx(path: str | Path) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [_cell(c) for c in next(it, [])]
        rows = []
        for r in it:
            if not any(c is not None and str(c).strip() for c in r):
                continue   # Leerzeilen überspringen
            rows.append({header[i]: _cell(v) for i, v in enumerate(r) if i < len(header)})
        return header, rows
    finally:
        wb.close()


def _read_xls(path: str | Path) -> tuple[list[str], list[dict]]:
    try:
        import xlrd
    except ModuleNotFoundError as e:
        raise ValueError(".xls-Dateien benötigen das Paket 'xlrd' (pip install xlrd) — "
                         "oder die Datei als .xlsx/.csv speichern") from e
    book = xlrd.open_workbook(path)
    sh = book.sheet_by_index(0)
    if sh.nrows == 0:
        return [], []
    header = [_cell(sh.cell_value(0, c)) for c in range(sh.ncols)]
    rows = [{header[c]: _cell(sh.cell_value(r, c)) for c in range(sh.ncols)}
            for r in range(1, sh.nrows)]
    return header, rows


def read_table(path: str | Path) -> tuple[list[str], list[dict]]:
    """(header, rows) aus CSV/XLSX/XLS. rows sind dicts {Spaltenname: String}.
    Wirft ValueError bei unbekannter Endung oder fehlender .xls-Abhängigkeit."""
    ext = Path(path).suffix.lower()
    if ext == ".csv":
        return _read_csv(path)
    if ext == ".xlsx":
        return _read_xlsx(path)
    if ext == ".xls":
        return _read_xls(path)
    raise ValueError(f"Format nicht unterstützt: {ext or '(keine Endung)'} "
                     f"(erlaubt: {', '.join(SUPPORTED_EXTENSIONS)})")


def guess_mapping(header: list[str]) -> dict:
    """Spalten -> unsere Felder. Zwei Durchgänge: erst exakte Treffer auf einen bekannten Spaltennamen
    (normalisiert, deshalb egal ob „Track-Nr.", „track number" oder „TRACK: Number"), dann für die noch
    offenen Felder zusammengesetzte Köpfe wie „CODE: GEMA WORK NUMBER" oder „German Label Code".
    Jede Spalte wird höchstens einmal vergeben."""
    by_norm: dict[str, str] = {}
    for h in header:
        k = _norm(h)
        if k and k not in by_norm:
            by_norm[k] = h
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for field, aliases in _ALIASES.items():
        for a in aliases:
            col = by_norm.get(_norm(a))
            if col and col not in used:
                mapping[field] = col
                used.add(col)
                break
    for field in _SUBSTRING_OK:
        if field in mapping:
            continue
        for a in sorted({_norm(x) for x in _ALIASES[field]}, key=len, reverse=True):
            if len(a) < 4 or a in _SUBSTRING_TOO_GENERIC:
                continue
            hit = next((col for k, col in by_norm.items() if a in k and col not in used), None)
            if hit:
                mapping[field] = hit
                used.add(hit)
                break
    return mapping


def library_delivery_kind(path: str | Path) -> str:
    """'upm' wenn die Datei eine Track-Lieferung einer Library ist, sonst ''. Solche Dateien gehen NICHT
    durch das Spalten-Mapping, sondern über den Lieferungs-Weg (verifiziert, mit Dateinamen-Tabelle für die
    Cue-Zuordnung); Library-Metadaten gelten als verifiziert."""
    from . import librarytracks
    if Path(path).suffix.lower() != ".csv":
        return ""
    try:
        with open(path, newline="", encoding="utf-8-sig") as f:
            first = f.readline()
            if not first.lstrip().lower().startswith("sep="):
                f.seek(0)
            names = csv.DictReader(f).fieldnames
    except Exception:
        return ""
    return "upm" if librarytracks.looks_like_upm_csv(names) else ""


def preview_library_delivery(path: str | Path, code2name) -> dict:
    """Trockenlauf für eine Library-Lieferung: {kind, library, rows, sample, labels, unknown_codes, corrected}.
    labels {code: {label, n}}; Zeilen ohne Code unter „—“ mit names {gelieferter Labelname: n}.
    code2name am besten labelcheck.GvlMaster — dann gilt der GVL-Code statt eines fremden (s. to_song_rows)."""
    from . import librarytracks
    kind = library_delivery_kind(path)
    if not kind:
        return {"error": "Keine erkannte Library-Lieferung."}
    rows = librarytracks.read_upm_csv(path)
    songs, rep = librarytracks.to_song_rows(rows, code2name)
    labels: dict[str, dict] = {}
    for sgl in songs:
        e = labels.setdefault(sgl["labelcode"] or "—", {"label": sgl["labelname"], "n": 0})
        e["n"] += 1
        if not sgl["labelcode"]:      # ohne Code die gelieferten Labelnamen zählen — sonst nur „— · LC—“
            names = e.setdefault("names", {})
            names[sgl["sublibrary"]] = names.get(sgl["sublibrary"], 0) + 1
    return {"kind": kind, "library": kind, "rows": len(songs), "labels": labels,
            "unknown_codes": rep["unknown_codes"], "corrected": rep["corrected"], "no_code": rep["no_code"],
            "by_name": rep["by_name"],
            "sample": [{k: sgl.get(k, "") for k in ("file", "interpret", "titel", "komponist", "verlag",
                                                    "labelname", "labelcode", "album", "isrc")}
                       for sgl in songs[:5]]}


def import_library_delivery(con, path: str | Path, code2name, *, date_verified: str) -> dict:
    """Library-Lieferung einpflegen (verifiziert, mit library_files für die Cue-Zuordnung). Die fremden Codes
    (unknown_codes) muss der Aufrufer als manuelle Labels anlegen — die Label-Datenbank liegt nicht hier."""
    from . import librarytracks
    kind = library_delivery_kind(path)
    if not kind:
        return {"error": "Keine erkannte Library-Lieferung."}
    rows = librarytracks.read_upm_csv(path)
    songs, rep = librarytracks.to_song_rows(rows, code2name)
    res = songsdb.import_library_tracks(con, songs, library=kind, date_verified=date_verified)
    return {"library": kind, **res, "rows": rep["rows"], "unknown_codes": rep["unknown_codes"],
            "corrected": rep["corrected"], "renamed": rep["renamed"], "no_code": rep["no_code"],
            "by_name": rep["by_name"]}


def find_deliveries(folder: str | Path, *, max_depth: int = 4, limit: int = 500) -> list[Path]:
    """Alle Library-Lieferungs-CSVs unter `folder` (rekursiv, begrenzte Tiefe). Manche Libraries legen bei JEDEM Download
    eine `music_metadata.csv` neben die Dateien — sie sammeln sich also in den Download-Ordnern an, statt einmal
    komplett zu kommen. Deshalb Ordner scannen statt einzeln hochladen."""
    root = Path(folder)
    if not root.is_dir():
        return []
    out: list[Path] = []
    for p in sorted(root.rglob("*.csv")):
        if len(p.relative_to(root).parts) > max_depth:
            continue
        try:
            if p.stat().st_size > 20_000_000:        # Lieferungen sind klein; ein Riesen-CSV ist etwas anderes
                continue
        except OSError:
            continue
        if library_delivery_kind(p):
            out.append(p)
        if len(out) >= limit:
            break
    return out


def _merge_codes(into: dict, codes: dict) -> None:
    """unknown_codes/corrected/by_name mehrerer Lieferungen zusammenzählen."""
    for code, d in codes.items():
        e = into.setdefault(code, {**d, "n": 0})
        e["n"] += d["n"]


def preview_delivery_folder(folder: str | Path, code2name) -> dict:
    """Trockenlauf über einen Ordner: je gefundener Lieferung Datei, Zeilen und Labels; dazu die Summe samt
    fremden (unknown_codes) und ersetzten Codes (corrected) — die Vorschau muss sie zeigen."""
    files = find_deliveries(folder)
    items, total, labels, unknown, corrected, by_name, no_code = [], 0, {}, {}, {}, {}, 0
    for f in files:
        pv = preview_library_delivery(f, code2name)
        if "error" in pv:
            items.append({"file": str(f), "error": pv["error"]})
            continue
        total += pv["rows"]
        for code, d in pv["labels"].items():
            e = labels.setdefault(code, {"label": d["label"], "n": 0})
            e["n"] += d["n"]
            if "names" in d:
                names = e.setdefault("names", {})
                for name, n in d["names"].items():
                    names[name] = names.get(name, 0) + n
        _merge_codes(unknown, pv["unknown_codes"])
        _merge_codes(corrected, pv["corrected"])
        _merge_codes(by_name, pv["by_name"])
        no_code += pv["no_code"]
        items.append({"file": str(f), "name": f.name, "parent": f.parent.name, "rows": pv["rows"],
                      "labels": pv["labels"], "unknown_codes": pv["unknown_codes"], "corrected": pv["corrected"],
                      "by_name": pv["by_name"]})
    return {"folder": str(folder), "found": len(files), "rows": total, "labels": labels, "items": items,
            "unknown_codes": unknown, "corrected": corrected, "by_name": by_name, "no_code": no_code}


def import_delivery_folder(con, folder: str | Path, code2name, *, date_verified: str) -> dict:
    """Alle Lieferungen eines Ordners einpflegen. Idempotent: ein zweiter Lauf zählt als `updated`,
    nicht als neu. unknown_codes je Datei, damit der Aufrufer die manuellen Labels mit Quelle anlegt."""
    res = {"folder": str(folder), "files": 0, "inserted": 0, "updated": 0, "file_rows": 0, "items": []}
    for f in find_deliveries(folder):
        r = import_library_delivery(con, f, code2name, date_verified=date_verified)
        if "error" in r:
            res["items"].append({"file": str(f), "error": r["error"]})
            continue
        res["files"] += 1
        res["inserted"] += r["inserted"]
        res["updated"] += r["updated"]
        res["file_rows"] += r["files"]
        res["items"].append({"file": str(f), "name": f.name, "library": r["library"], "inserted": r["inserted"],
                             "updated": r["updated"], "unknown_codes": r["unknown_codes"], "corrected": r["corrected"],
                             "by_name": r["by_name"]})
    return res


def _values(row: dict, mapping: dict) -> dict:
    """Zeile -> {unser Feld: Wert}, getrimmt. Labelcodes kommen als „LC 48731" oder „48731" und werden
    auf die Repo-Schreibweise normalisiert (führende Nullen, kein Leerzeichen)."""
    from .catalog import clean_writer_names
    from .librarytracks import flip_person_list
    from .normalize import strict_labelcode
    out = {f: (row.get(c, "") or "").strip() for f, c in mapping.items()}
    if out.get("labelcode"):
        digits = strict_labelcode(out["labelcode"])
        # nicht lesbar ('LC 24-584', zwei Codes): wie geliefert stehen lassen, _check_label markiert die Zeile
        # zur Prüfung — nie still auf den ersten Ziffernblock kürzen
        out["labelcode"] = ("LC" + digits) if digits else (out["labelcode"] if digits is None else "")
    # Urheber- und Verlagsspalten tragen in Cue-Sheets Verwertungsgesellschaft, IPI und Anteil mit
    # („Erika Testerin (PRS: 100000001, 100%)") — raus damit, Mehrfachnennungen auf „; " vereinheitlicht.
    # Personen im Format „Nachname, Vorname" werden vorher gedreht, sonst teilt der Komma-Split sie.
    for f in ("komponist", "texter", "verlag"):
        if out.get(f):
            out[f] = clean_writer_names(out[f] if f == "verlag" else flip_person_list(out[f]))
    # „1.0" aus Tabellen, die Zahlen als Fließkomma exportieren
    if out.get("track_number", "").endswith(".0") and out["track_number"][:-2].isdigit():
        out["track_number"] = out["track_number"][:-2]
    return out


def _check_label(v: dict, master) -> bool:
    """Labelcode einer fremden Liste gegen den GVL-Master (core/labelcheck). Passen Code und Name, steht der
    Name danach in GVL-Schreibweise da; sonst bleibt beides wie geliefert (der Name in GVL-Schreibweise, wenn der
    Master ihn kennt) und die Zeile wird zur Prüfung markiert. Rückgabe: True = zur Prüfung."""
    from .labelcheck import check_label
    from .labelmatch import norm
    if master is None or not v.get("labelcode"):
        return False
    chk = check_label(v["labelcode"], v.get("labelname", ""), master)
    if chk["status"] == "ok":
        v["labelname"] = chk["label"]
        return False
    for code in master.codes_for_name(v.get("labelname", ""))[:1]:
        v["labelname"] = next((n for n in master.names(code) if norm(n) == norm(v["labelname"])), v["labelname"])
    return True


def preview_file(path: str | Path, *, master=None) -> dict:
    """Trockenlauf: erkannte Spalten + Mapping + Zeilenzahl + bis zu 5 Beispielzeilen.
    Schreibt nichts. Gibt {fieldnames, mapping, rows, required_ok, sample, label_review, label_review_codes}
    oder {error,...}. master (labelcheck.GvlMaster): Labelcodes gegen den GVL-Master prüfen."""
    try:
        header, rows = read_table(path)
    except ValueError as e:
        return {"error": str(e)}
    mp = guess_mapping(header)
    vals = [_values(row, mp) for row in rows]
    review = [v["labelcode"] for v in vals if _check_label(v, master)]
    from_comp = sum(1 for v in vals if not v.get("interpret") and v.get("komponist"))
    usable = sum(1 for v in vals if v.get("titel") and (v.get("interpret") or v.get("komponist")))
    return {"fieldnames": header, "mapping": mp, "rows": len(rows), "fields": list(_ALIASES),
            "required_ok": usable > 0, "usable": usable, "artist_from_composer": from_comp,
            "label_review": len(review), "label_review_codes": sorted(set(review))[:10],
            "sample": [{f: v.get(f, "") for f in mp} for v in vals[:5]]}


def import_file(con, path: str | Path, *, quelle: str = "", mapping: dict | None = None, master=None) -> dict:
    """CSV/XLSX/XLS importieren + per upsert in songs.db mergen (verifizierte Felder
    bleiben geschützt). Gibt {inserted, updated, skipped, mapping, label_review} oder {error,...}.
    master (labelcheck.GvlMaster): Zeilen, deren Labelcode nicht zum Master passt, werden nicht als verifiziert
    übernommen, sondern mit review_status 'review' markiert — nur wenn ihr Code auch geschrieben wird."""
    try:
        header, rows = read_table(path)
    except ValueError as e:
        return {"error": str(e)}
    mp = {f: c for f, c in (mapping or guess_mapping(header)).items() if c in header}
    if "titel" not in mp or not ("interpret" in mp or "komponist" in mp):
        return {"error": "Pflichtspalten nicht erkannt: Titel, dazu Interpret oder Komponist",
                "fieldnames": header}
    ins = upd = skip = from_comp = review = 0
    for row in rows:
        v = _values(row, mp)
        flag = _check_label(v, master)
        if not v.get("interpret") and v.get("komponist"):
            # Steht bei Production Music kein Interpret, IST der Komponist der Interpret — so läuft
            # die GEMA-Meldung. Viele Listenformate führen gar keine Interpret-Spalte.
            v["interpret"] = v["komponist"]
            from_comp += 1
        if not v.get("interpret") or not v.get("titel"):
            skip += 1
            continue
        if flag:
            # Nur markieren, wenn dieser Code auch geschrieben wird (labelcode ist PROTECT): eine Zeile, die schon
            # einen Code trägt, wird von einer fremden Liste nicht herabgestuft
            prev = songsdb.find(con, v["interpret"], v["titel"])
            flag = not (prev and (prev.get("labelcode") or "").strip())
            review += flag
        r = songsdb.upsert(con, interpret=v["interpret"], titel=v["titel"],
                           komponist=v.get("komponist", ""), texter=v.get("texter", ""),
                           verlag=v.get("verlag", ""),
                           labelname=v.get("labelname", ""), labelcode=v.get("labelcode", ""),
                           album=v.get("album", ""), track_number=v.get("track_number", ""),
                           isrc=v.get("isrc", ""), iswc=v.get("iswc", ""),
                           gema_werknummer=v.get("gema_werknummer", ""), quelle=quelle,
                           review_status="review" if flag else "")
        ins += r == "inserted"
        upd += r == "updated"
    con.commit()
    return {"inserted": ins, "updated": upd, "skipped": skip, "mapping": mp,
            "artist_from_composer": from_comp, "label_review": review}


def import_csv(con, path: str | Path, *, quelle: str = "", mapping: dict | None = None) -> dict:
    """Rückwärtskompatibler Alias (Ordner-CLI + Alt-Tests). Delegiert an import_file."""
    return import_file(con, path, quelle=quelle, mapping=mapping)

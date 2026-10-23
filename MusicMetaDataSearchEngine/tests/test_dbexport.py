import csv

from core import dbexport, ingest, songsdb

ROWS = [
    {"interpret": "AURORA", "titel": "Runaway", "komponist": "X; Y", "texter": "Z",
     "album": "All My Demons", "track_number": "2", "labelname": "Decca", "labelcode": "00171",
     "isrc": "GBUM71000001", "iswc": "T-000.000.001-0", "gema_werknummer": "999"},
    {"interpret": "Bon Iver", "titel": "Holocene", "album": "Bon Iver", "track_number": "5",
     "labelname": "Jagjaguwar", "labelcode": "LC00382"},
]


def test_export_csv_roundtrips_through_ingest(tmp_path):
    p = tmp_path / "out.csv"
    dbexport.export_csv(ROWS, p)
    # Header + Werte lesbar
    with open(p, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    assert rows[0][0] == "Interpret" and "Track-Nr." in rows[0]
    assert rows[1][0] == "AURORA"
    # labelcode wird einheitlich mit LC-Präfix geschrieben
    lc_col = rows[0].index("LC")
    assert rows[1][lc_col] == "LC00171"

    # Roundtrip: der Export muss sich sauber reimportieren lassen
    con = songsdb.open(tmp_path / "songs.db")
    res = ingest.import_file(con, p)
    assert res["inserted"] == 2
    hit = songsdb.find(con, "AURORA", "Runaway")
    assert hit["album"] == "All My Demons" and hit["track_number"] == "2"
    assert hit["labelcode"] == "LC00171" and hit["isrc"] == "GBUM71000001"
    con.close()


def test_export_xlsx_readable(tmp_path):
    p = tmp_path / "out.xlsx"
    dbexport.export_xlsx(ROWS, p)
    from openpyxl import load_workbook
    ws = load_workbook(p).active
    data = list(ws.iter_rows(values_only=True))
    assert data[0][0] == "Interpret"
    assert data[1][0] == "AURORA"
    # Reimport über ingest bestätigt den Roundtrip auch für XLSX
    con = songsdb.open(tmp_path / "songs.db")
    assert ingest.import_file(con, p)["inserted"] == 2
    con.close()


def test_export_pdf_is_valid_pdf(tmp_path):
    p = tmp_path / "out.pdf"
    dbexport.export_pdf(ROWS, p)
    assert p.read_bytes()[:5] == b"%PDF-"
    import fitz
    doc = fitz.open(p)
    assert doc.page_count >= 1
    text = doc[0].get_text()
    assert "AURORA" in text and "CueCore" in text
    doc.close()


def test_export_pdf_handles_empty(tmp_path):
    p = tmp_path / "empty.pdf"
    dbexport.export_pdf([], p)
    assert p.read_bytes()[:5] == b"%PDF-"


def test_export_has_verlag_column_and_roundtrips(tmp_path):
    rows = [dict(ROWS[0], verlag="Universal Music Publishing")]
    p = tmp_path / "out.csv"
    dbexport.export_csv(rows, p)
    with open(p, encoding="utf-8-sig", newline="") as f:
        head, first = list(csv.reader(f))[:2]
    assert head.index("Verlag") == head.index("Texter") + 1            # direkt neben den Autoren, vor dem Label
    assert first[head.index("Verlag")] == "Universal Music Publishing"
    con = songsdb.open(tmp_path / "songs.db")
    ingest.import_file(con, p)
    assert songsdb.find(con, "AURORA", "Runaway")["verlag"] == "Universal Music Publishing"
    con.close()


def test_labelcode_export_is_lc_digits_and_keeps_unreadable_values():
    """Jeder Export: 'LC' + Ziffern ohne Leerzeichen (nicht 'LC121' oder 'LC 24 584')."""
    assert dbexport._val({"labelcode": "121"}, "labelcode") == "LC00121"
    assert dbexport._val({"labelcode": "LC 24 584"}, "labelcode") == "LC24584"
    assert dbexport._val({"labelcode": "LC 01485"}, "labelcode") == "LC01485"
    assert dbexport._val({"labelcode": "LC104994"}, "labelcode") == "LC104994"
    assert dbexport._val({"labelcode": ""}, "labelcode") == ""
    assert dbexport._val({"labelcode": "unbekannt"}, "labelcode") == "unbekannt"


def test_export_lc_column_is_always_lc_plus_digits(tmp_path):
    """Auch im Datenbank-Export: 'LC' + Ziffern mit führenden Nullen, ohne Leerzeichen (nicht 'LC171' oder
    'LC 21110' roh)."""
    rows = [dict(ROWS[1], labelcode="171"), dict(ROWS[1], titel="B", labelcode="LC 21110"),
            dict(ROWS[1], titel="C", labelcode="")]
    p = tmp_path / "out.csv"
    dbexport.export_csv(rows, p)
    with open(p, encoding="utf-8-sig", newline="") as f:
        data = list(csv.reader(f))
    lc = data[0].index("LC")
    assert [r[lc] for r in data[1:]] == ["LC00171", "LC21110", ""]


def _many(n):
    return [{"interpret": f"Interpret {i:05d}", "titel": f"Titel Nummer {i:05d}", "komponist": "Pierre-André Athané; X",
             "verlag": "Frederic Leibovitz Editeur", "album": "Little French Fancies", "track_number": str(i % 30),
             "labelname": "CMG", "labelcode": "LC10347", "isrc": f"FR26F03{i:05d}"} for i in range(n)]


def test_export_pdf_time_grows_linearly(tmp_path):
    """Eine einzige Story-Tabelle über alle Zeilen wüchse überproportional. Viermal so viele Zeilen dürfen
    nur etwa viermal so lange brauchen (nicht den Faktor 11)."""
    import time
    t = []
    for n in (1000, 4000):
        start = time.perf_counter()
        dbexport.export_pdf(_many(n), tmp_path / f"{n}.pdf")
        t.append(time.perf_counter() - start)
    assert t[1] < 6 * t[0], t


def test_export_pdf_multi_page_keeps_rows_and_columns(tmp_path):
    """Gestückelte Ausgabe: jede Zeile genau einmal, Kopf auf Seite 1, die Spalten stehen auf jeder Seite
    an derselben Stelle, und ein Wort ohne Trennstelle schiebt nichts über den Rand."""
    import fitz
    rows = _many(450)
    rows[300]["titel"] = "x" * 150
    p = tmp_path / "many.pdf"
    dbexport.export_pdf(rows, p)
    doc = fitz.open(p)
    assert doc.page_count > 5
    text = "".join("".join(pg.get_text().split()) for pg in doc)
    assert all(text.count(f"Interpret{i:05d}") == 1 for i in range(450))
    assert "Track-Nr." in doc[0].get_text() and "450 Einträge" in doc[0].get_text()
    lc_x = {round(w[0], 1) for pg in doc for w in pg.get_text("words") if w[4] == "LC10347"}
    assert len(lc_x) == 1, lc_x
    limit = doc[0].rect.x1 - 36 + 0.5
    assert max(w[2] for pg in doc for w in pg.get_text("words")) <= limit
    assert "x" * 150 in text
    doc.close()


def test_export_pdf_wiederholt_den_tabellenkopf_auf_jeder_seite(tmp_path):
    """Der Tabellenkopf stand nur auf Seite 1 — ab Seite 2 waren zwölf Spalten ohne Überschrift zu lesen.
    Jetzt steht er auf jeder Seite einmal, oben über den Zeilen, an derselben Stelle; Titel und Anzahl nur
    auf Seite 1; keine Zeile geht verloren oder doppelt."""
    import fitz
    rows = _many(450)
    for r in rows:
        r["interpret"] = r["interpret"].replace("Interpret ", "Zonk")        # kein Datenwort gleich einer Überschrift
    p = tmp_path / "kopf.pdf"
    dbexport.export_pdf(rows, p)
    doc = fitz.open(p)
    assert doc.page_count > 5
    heads = ("Interpret", "Komponist", "Track-Nr.", "GEMA-Werknr.")       # stehen in keiner Datenzeile
    pos = set()
    for i, pg in enumerate(doc):
        words = pg.get_text("words")
        text = pg.get_text()
        for head in (*heads, "LC"):
            assert sum(1 for w in words if w[4] == head) == 1, (i, head)
        head_y = max(w[3] for w in words if w[4] in heads)
        assert head_y < min(w[1] for w in words if w[4].startswith("FR-26F-03")), i   # Kopf über den Zeilen (ISRC im PDF mit Bindestrichen)
        pos.add(tuple(round(w[0], 1) for w in words if w[4] in ("Interpret", "LC", "GEMA-Werknr.")))
        assert ("450 Einträge" in text) == (i == 0)
    assert len(pos) == 1, pos
    text = "".join("".join(pg.get_text().split()) for pg in doc)
    assert all(text.count(f"Zonk{i:05d}") == 1 for i in range(450))
    doc.close()
